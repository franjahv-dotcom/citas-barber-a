from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math
ROOT=Path.cwd(); OUT=ROOT/'output/documentos'; TMP=ROOT/'tmp/documentacion'; AS=TMP/'figuras'
OUT.mkdir(parents=True,exist_ok=True); AS.mkdir(parents=True,exist_ok=True)
FONT='C:/Windows/Fonts/arial.ttf'; BOLD='C:/Windows/Fonts/arialbd.ttf'
def diagram(name,size,boxes,arrows):
 im=Image.new('RGB',size,'white'); dr=ImageDraw.Draw(im)
 for coords,label in arrows:
  dr.line(coords,fill='#687783',width=4)
  (x1,y1),(x2,y2)=coords[-2:];a=math.atan2(y2-y1,x2-x1)
  dr.polygon([(x2,y2),(x2-15*math.cos(a-.45),y2-15*math.sin(a-.45)),(x2-15*math.cos(a+.45),y2-15*math.sin(a+.45))],fill='#687783')
  if label:
   x,y=coords[len(coords)//2];dr.text((x+10,y-30),label,font=ImageFont.truetype(FONT,22),fill='#374151')
 for (x,y,w,h),title,body in boxes:
  dr.rounded_rectangle((x,y,x+w,y+h),radius=12,fill='#edf2f6',outline='#607383',width=3)
  dr.text((x+w/2,y+25),title,font=ImageFont.truetype(BOLD,27),fill='black',anchor='mt')
  for i,line in enumerate(body):dr.text((x+w/2,y+72+i*33),line,font=ImageFont.truetype(FONT,23),fill='#20252a',anchor='mt')
 im.save(AS/(name+'.png'))
diagram('arquitectura',(1500,390),[
 ((20,65,390,220),'Navegador',['React y Tailwind CSS','Cliente o administrador','Formularios y agenda']),
 ((555,65,390,220),'Servidor',['Node.js y Express','Sesiones y validaciones','Disponibilidad y transacciones']),
 ((1090,65,390,220),'Base de datos',['MySQL 8.0 en Docker','Citas e historial','Jornadas y bloqueos'])],
 [([(410,150),(555,150)],''), ([(945,150),(1090,150)],'')])
diagram('reserva',(1450,660),[
 ((20,25,410,190),'1 Elegir',['Servicio, barbero y fecha','Consultar horarios','Registrar nombre y teléfono']),
 ((520,25,410,190),'2 Validar',['Iniciar transacción','Bloquear fila de agenda','Revisar disponibilidad']),
 ((1010,25,410,190),'3 Decidir',['Jornada y duración completas','Sin descanso ni bloqueo','Sin cruce con cita activa']),
 ((1010,385,410,190),'4 Guardar',['Relacionar al cliente','Guardar cita y movimiento','Confirmar transacción']),
 ((520,385,410,190),'Respuesta al cliente',['Mostrar número de reserva','Fecha y hora confirmadas','Liberar el bloqueo']),
 ((20,385,410,190),'Si no es válido',['Revertir transacción','Avisar del conflicto','Conservar datos anteriores'])],
 [([(430,120),(520,120)],''), ([(930,120),(1010,120)],''), ([(1215,215),(1215,385)],''), ([(1010,480),(930,480)],''), ([(725,215),(725,295),(225,295),(225,385)],'')])
diagram('datos',(1480,760),[
 ((20,20,350,160),'Clientes',['id, nombre y teléfono','Teléfono no único']),
 ((565,20,350,160),'Citas',['cliente, servicio y barbero','inicio, fin y estado']),
 ((1110,20,350,160),'Servicios',['nombre y descripción','precio, duración y activo']),
 ((565,325,350,160),'Barberos',['nombre y activo','servicios que realiza']),
 ((20,570,350,160),'Jornadas',['barbero, día semanal','hora inicial y final']),
 ((1110,570,350,160),'Bloqueos',['barbero, inicio y fin','motivo de ausencia']),
 ((20,325,350,160),'Movimientos',['cita, acción y detalle','registro del cambio']),
 ((1110,325,350,160),'Barbero y servicio',['tabla de asociación','barber_services'])],
 [([(370,95),(565,95)],''), ([(1110,95),(915,95)],''), ([(740,325),(740,180)],''), ([(565,120),(465,120),(465,400),(370,400)],''), ([(565,410),(465,410),(465,650),(370,650)],''), ([(915,445),(1005,445),(1005,650),(1110,650)],''), ([(915,365),(1110,365)],''), ([(1290,180),(1290,325)],'')])
D=Document(); sec=D.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.66);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.footer_distance=Inches(.25)
styles=D.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.space_after=Pt(7);styles['Normal'].paragraph_format.line_spacing=1.12
for name,size in [('Title',30),('Subtitle',15),('Heading 1',19),('Heading 2',12.5),('Heading 3',11)]:
 styles[name].font.size=Pt(size);styles[name].paragraph_format.space_before=Pt(10);styles[name].paragraph_format.space_after=Pt(8)
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.italic=True
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run('Mala Vida Barbers  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
D.core_properties.title='Documentación de la segunda entrega del sistema de citas de barbería';D.core_properties.author='Emmanuel Perea'
page=0
def new(title=None):
 global page
 if page:D.add_page_break()
 page+=1
 if title:D.add_heading(title,1)
def p(t,bold=False):
 x=D.add_paragraph();r=x.add_run(t);r.bold=bold;return x
def h(t):D.add_heading(t,2)
def bullet(t):
 x=D.add_paragraph(t,style='List Bullet');x.paragraph_format.space_after=Pt(4);return x
def cap(t):D.add_paragraph(t,'Caption')
def pic(file,width=6.7):
 x=D.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(4)
 x.add_run().add_picture(str(file),width=Inches(width));x._p.xpath('.//wp:docPr')[0].set('descr',Path(file).stem.replace('_',' '))
def table(headers,rows,widths,font=9):
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,text in enumerate(headers):t.rows[0].cells[i].text=text
 for row in rows:
  for i,value in enumerate(row):t.add_row().cells[i].text=str(value) if False else ''
 # Fill rows without creating partial rows.
 for row in list(t.rows)[1:]:t._tbl.remove(row._tr)
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'253B4B' if ri==0 else ('F1F4F6' if ri%2==0 else 'FFFFFF'));cp.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   cp.append(borders);margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
   cp.append(margins)
   for pa in c.paragraphs:
    pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.line_spacing=1.03
    if ci==0 and widths[0]<.8:pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in pa.runs:run.font.size=Pt(font);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
 D.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
new()
p('INGE00031  |  INGENIERÍA DE SOFTWARE',True)
D.add_paragraph().paragraph_format.space_after=Pt(35)
D.add_paragraph('Sistema web de gestión y control de citas para una barbería','Title')
D.add_paragraph('Diseño, desarrollo, implementación y evaluación del prototipo','Subtitle')
p('Mala Vida Barbers',True)
D.add_paragraph().paragraph_format.space_after=Pt(25)
p('Alumno  Emmanuel Perea')
p('Docente  ____________________________________')
p('Institución y grupo  ____________________________')
p('Fecha  23 de septiembre de 2026')
D.add_paragraph().paragraph_format.space_after=Pt(25)
h('Propósito del documento')
p('Presentar la segunda etapa del proyecto y reunir los materiales necesarios para elaborar el informe, la presentación y la demostración. El desarrollo toma como punto de partida los requerimientos de la primera entrega, que se conserva sin modificaciones.')
h('Resultado principal')
p('El prototipo integra reservas, administración de la agenda, jornadas, bloqueos e historial. La ejecución de verificación registrada reporta 27 resultados aprobados y ningún fallo; ese total incluye 26 casos y el resultado del grupo de integración. La evaluación corresponde a un entorno local controlado.')
h('Organización para la entrega')
p('Las páginas 1 a 10 forman el informe principal. A partir de la página 11 se incluyen anexos para respaldar las pruebas, preparar las diapositivas y repartir el trabajo de exposición. Los anexos pueden entregarse por separado si el docente limita la extensión del informe.')
new('Índice e introducción')
table(['Contenido del informe','Página'],[['Portada y propósito','1'],['Índice e introducción','2'],['Diseño de la arquitectura','3'],['Diseño de procesos y reglas','4'],['Diseño de la base de datos','5'],['Diseño de la interfaz','6'],['Desarrollo e implementación','7'],['Pruebas y resultados','8'],['Evaluación de calidad','9'],['Conclusiones y referencias','10']],[5.75,1])
h('Introducción')
p('Una barbería necesita conocer qué trabajador puede atender, cuánto dura cada servicio y qué espacios siguen libres. Si esa información se reparte entre mensajes y anotaciones, es posible aceptar citas incompatibles o perder el seguimiento de un cambio. La primera entrega definió ese problema, sus actores y los requisitos del sistema.')
p('En esta segunda etapa se implementó un prototipo web que centraliza la información. El cliente consulta servicios y horarios y reserva sin crear una cuenta. El administrador inicia sesión para consultar y modificar la agenda, mantener los catálogos y resolver cancelaciones, reprogramaciones y ausencias.')
p('El objetivo de esta fase es comprobar que el diseño se traduce en operaciones consistentes y verificables. Por ello se documentan la arquitectura, los procesos, la base de datos, la interfaz y las pruebas. El criterio principal es conservar una agenda válida aun cuando varias personas soliciten el mismo horario.')
h('Alcance y condiciones')
p('El proyecto contempla servicios, barberos, clientes, citas, jornadas, descansos, bloqueos y acceso administrativo. Se mantienen fuera del alcance los pagos en línea, el inventario, la nómina, los recordatorios automáticos y las cuentas individuales de barberos. Los horarios y las políticas iniciales son supuestos del prototipo académico; no se presentan como resultados de entrevistas con un negocio real.')
new('Diseño de la arquitectura')
h('Separación de responsabilidades')
p('El sistema utiliza una arquitectura cliente servidor. React presenta los formularios y la agenda; Node.js con Express recibe las solicitudes y aplica las reglas; MySQL conserva los datos. El navegador se comunica con la API y no accede directamente a la base de datos.')
pic(AS/'arquitectura.png');cap('Figura 1. Arquitectura implementada. La interfaz usa HTTP y JSON; el servidor consulta MySQL mediante SQL parametrizado.')
table(['Componente','Responsabilidad','Elementos principales'],[['Interfaz','Capturar datos y mostrar respuestas','Reserva, agenda, catálogos, jornadas, bloqueos y clientes'],['API','Autorizar y validar operaciones','Sesiones, rutas públicas y administrativas, reglas de disponibilidad'],['Persistencia','Guardar relaciones e historial','MySQL, claves foráneas, transacciones e índices']],[1.05,2.25,3.45])
h('Organización de los módulos')
p('El servidor separa validaciones generales, autenticación, acceso a datos, reglas de agenda, inicialización y rutas. En la interfaz, cada función administrativa tiene un componente propio y comparte utilidades para solicitudes, mensajes, cuadros de diálogo y estados de carga. Esta organización facilita revisar una regla sin repetirla en cada pantalla.')
h('Acceso y comunicación')
p('La consulta pública devuelve servicios, barberos e intervalos disponibles, sin exponer nombres ni teléfonos de otros clientes. La agenda y las operaciones de administración requieren una sesión vigente. Las contraseñas se almacenan mediante scrypt con una sal aleatoria; el navegador recibe una cookie de sesión HttpOnly y SameSite Strict. La documentación de Node.js fundamenta el uso de scrypt como función de derivación de contraseñas [3].')
p('En la ejecución local, el servidor entrega la interfaz y la API en la misma dirección. Para desarrollo, Vite redirige las solicitudes de la API al servidor. Esta separación permite trabajar en la interfaz sin cambiar las reglas de negocio.')
new('Diseño de procesos y reglas')
p('La disponibilidad mostrada al cliente es orientativa hasta confirmar. El servidor repite la validación dentro de una transacción, de modo que consultar un horario libre no garantiza que continúe libre cuando se intente guardar.')
pic(AS/'reserva.png');cap('Figura 2. Reserva y respuesta ante un rechazo. El camino inferior izquierdo corresponde a una validación fallida.')
h('Regla de conflicto')
p('Existe un cruce si inicio_nueva < fin_existente y fin_nueva > inicio_existente para el mismo barbero y una cita pendiente o confirmada. Los límites pueden tocarse: una cita que termina a las 10:40 permite iniciar otra a las 10:40. El intervalo completo debe caber en la jornada y excluir descansos y bloqueos.')
h('Control de solicitudes simultáneas')
p('Las escrituras de agenda y configuración obtienen un bloqueo sobre una fila de control mediante SELECT FOR UPDATE. Solo después se revisa y guarda la cita. Las otras operaciones esperan, vuelven a consultar los datos y rechazan el horario si ya está ocupado. El bloqueo se conserva hasta confirmar o revertir la transacción, conforme al mecanismo de lecturas con bloqueo de InnoDB [2]. Esta decisión favorece la consistencia; también serializa las escrituras y limita la escalabilidad del prototipo.')
h('Cambios y estados')
p('La reprogramación valida el nuevo intervalo antes de modificar el registro. Si falla, la cita original permanece intacta. Cancelar conserva el historial y libera el espacio. Una cita confirmada puede completarse después de su hora de fin o marcarse como ausencia después de su inicio. El flujo público confirma directamente; el estado pendiente existe en el modelo y sus transiciones, pero no se ofrece un formulario público para crearlo. Este es un ajuste explícito frente al recorrido habitual previsto en la primera entrega.')
new('Diseño de la base de datos')
p('El modelo relacional separa los catálogos, la operación diaria y los registros de seguridad. Cada cita referencia un servicio, un barbero y, en los registros nuevos, un cliente. Las claves foráneas conservan esas relaciones; las bajas de catálogo se realizan mediante el estado activo.')
pic(AS/'datos.png');cap('Figura 3. Relaciones principales. Cada flecha va de la entidad referenciada a los registros asociados y representa una relación de uno a muchos.')
h('Decisiones de persistencia')
p('La cita guarda su inicio y fin calculados al reservar. Editar la duración de un servicio afecta las reservas nuevas, no desplaza los intervalos anteriores. El nombre y teléfono también se conservan en la cita como datos registrados en esa operación. El teléfono no es único: personas con nombres diferentes pueden compartirlo. Una coincidencia exacta de nombre y teléfono se reutiliza, por lo que identificar homónimos con el mismo número requiere revisión administrativa.')
p('Las jornadas admiten varios periodos por barbero y día. Los espacios entre periodos son descansos; periodos adyacentes se interpretan como una jornada continua. Los bloqueos tienen fecha inicial y final y pueden abarcar varios días. Si afectan citas existentes, se señalan para resolución sin eliminarlas.')
p('La estructura se completa con admins, admin_sessions, appointment_events, schedule_mutex y migrations. Los índices de agenda y búsqueda apoyan las consultas habituales. La inicialización migra la estructura anterior y relaciona sus citas con clientes sin borrar los registros existentes.')
new('Diseño de la interfaz')
p('La interfaz mantiene la identidad visual de Mala Vida Barbers con fondo oscuro, acentos dorados y controles de contraste visible. La pantalla pública reúne catálogo, selección de servicio y trabajador, fecha, horarios disponibles y datos de contacto. La confirmación muestra el número de reserva y el intervalo asignado.')
pic(ROOT/'backend/test-results/reserva-escritorio.png',6.1);cap('Figura 4. Formulario real durante la prueba automatizada. Los datos visibles son ficticios.')
h('Administración y prevención de errores')
p('El panel organiza agenda, clientes, servicios, barberos, jornadas, bloqueos y cuenta. Los filtros permiten localizar citas por fecha, trabajador, estado, nombre o teléfono. Cancelar y cambiar un estado requieren confirmación; los mensajes distinguen un conflicto de una operación guardada. Las citas afectadas por un bloqueo o una jornada modificada muestran una advertencia.')
p('La versión móvil distribuye formularios y tarjetas en una columna. Se verificó que no apareciera desplazamiento horizontal a 390 por 844 píxeles. Los controles tienen etiquetas y enfoque visible, y los estados se identifican con texto además del color. Estas comprobaciones no sustituyen una evaluación de accesibilidad ni una prueba de uso con participantes.')
new('Desarrollo e implementación')
h('Tecnologías y construcción')
p('Se implementó el cliente con React, Vite y Tailwind CSS; el servidor con JavaScript, Node.js y Express; y la persistencia con MySQL mediante mysql2. Docker conserva la base en un volumen. El desarrollo convierte los requisitos en módulos funcionales: reserva, administración, configuración e historial, seguidos de pruebas y correcciones.')
table(['Incremento funcional','Resultado implementado'],[['Reservas y disponibilidad','Cálculo de duración, horarios válidos, rechazo de cruces y protección concurrente'],['Administración','Autenticación, filtros, reprogramación, cancelación y cambios de estado'],['Configuración','Servicios, barberos, asignaciones, periodos semanales y bloqueos'],['Seguimiento y verificación','Clientes, movimientos, pruebas de API y recorridos en navegador']],[1.55,5.2])
p('La primera entrega propuso cuatro iteraciones de dos semanas y herramientas de apoyo como Trello, Figma y Visual Paradigm. La evidencia disponible aquí corresponde al código, diagramas de documentación y pruebas ejecutadas. No se atribuyen fechas de sprints, entrevistas ni entregables de esas herramientas que no estén respaldados por registros.')
h('Integración y conservación de datos')
p('La preparación inicial crea las tablas que faltan, adapta la estructura previa y registra la migración para evitar repetir la carga. Las consultas utilizan parámetros y validan identificadores, fechas, teléfono y duración. Las operaciones que afectan la agenda comparten el mismo control transaccional. Se verificó que repetir la inicialización no duplicara servicios ni jornadas.')
h('Configuración de demostración')
p('Se utiliza la hora de Ciudad de México. La jornada inicial es de lunes a sábado, de 10:00 a 14:00 y de 15:00 a 20:00; cada trabajador puede tener una configuración distinta. Las reservas se permiten con hasta 90 días de anticipación y se sugieren horarios cada 10 minutos. Los barberos nuevos necesitan una jornada y los servicios nuevos deben asignarse a quienes los realizan.')
p('El prototipo funciona localmente en http://127.0.0.1:3001. Las instrucciones de instalación están en el README y se resumen en el anexo D. El repositorio remoto sigue siendo un elemento pendiente de la entrega; una dirección local no lo sustituye.')
new('Pruebas y resultados')
h('Método de verificación')
p('La ejecución de respaldo se realizó el 23 de septiembre de 2026 con Node.js 24.19.0, MySQL 8.0.45 y Microsoft Edge 153.0.4234.48 sobre Windows. Se utilizó una base temporal independiente, datos ficticios y solicitudes HTTP reales. Al terminar se eliminó únicamente la base creada para las pruebas. El navegador se automatizó con Playwright 1.58.2.')
p('El registro informa 27 resultados aprobados: cuatro casos de reglas básicas, 22 casos dentro del grupo de integración y el resultado global de ese grupo. Por tanto, hay 26 casos ejecutables, no 27 escenarios independientes. El detalle completo y su correspondencia con los requisitos aparecen en los anexos A y B.')
table(['Comprobación','Resultado observado'],[['Solicitudes simultáneas','20 intentos sobre un mismo intervalo: 1 reserva creada y 19 rechazos por conflicto'],['Reprogramación inválida','Rechazo del cambio; inicio e identidad de la cita original conservados'],['Cancelación','Registro conservado y horario disponible nuevamente'],['Permisos','Consultas administrativas sin sesión rechazadas con HTTP 401'],['Bloqueos y jornadas','Se excluyen intervalos no disponibles y se señalan citas afectadas'],['Recorrido en navegador','Reserva, acceso, reprogramación, cancelación, historial y formularios administrativos completados'],['Resultado global','27 aprobados, 0 fallidos, 0 omitidos; duración aproximada de 14.32 segundos']],[2,4.75])
h('Rendimiento observado')
p('En una muestra de 20 consultas secuenciales de disponibilidad, sobre una base pequeña de pruebas y conexión local, la mediana fue 16.5 ms y el máximo 25.9 ms. El máximo quedó por debajo de la meta de 2 segundos planteada en RNF-04. La medición corresponde a la respuesta de la API, no al tiempo total de carga de una página ni a una prueba de carga de producción.')
h('Correcciones verificadas')
p('Las pruebas permitieron corregir la identificación accesible de los selectores y asegurar que los periodos contiguos se trataran como una jornada continua. La compilación y el análisis estático del frontend finalizaron sin errores ni advertencias en la revisión del desarrollo. Los resultados respaldan los casos probados, sin demostrar ausencia absoluta de defectos.')
new('Evaluación de calidad')
p('La evaluación toma ISO/IEC 25010:2023 como referencia general para organizar atributos y criterios medibles [4]. Se seleccionan aspectos pertinentes para el prototipo y se distinguen los resultados comprobados de las evaluaciones pendientes. No se afirma conformidad integral con la norma ni una certificación.')
table(['Aspecto','Criterio y evidencia','Valoración'],[['Adecuación funcional','Requisitos relacionados con módulos y pruebas; reservas y cambios verificables.','Núcleo implementado; ajuste del flujo pendiente explicado.'],['Fiabilidad e integridad','Concurrencia, rollback, límites de intervalo e historial comprobados.','Cumple los escenarios ejecutados; sin ensayo de fallos prolongados.'],['Eficiencia','20 consultas locales; mediana 16.5 ms y máximo 25.9 ms.','Cumple la meta local de 2 s; falta carga representativa.'],['Seguridad','Sesión requerida, scrypt, cookies protegidas, consultas parametrizadas y validación.','Controles básicos comprobados; sin auditoría de penetración.'],['Interacción y adaptación','Etiquetas, confirmaciones, avisos y vista móvil sin desbordamiento.','Verificación técnica favorable; falta evaluación con usuarios.'],['Mantenibilidad','Módulos separados, componentes compartidos y pruebas repetibles.','Estructura revisable; sin métrica de cobertura de código.'],['Compatibilidad del entorno','Edge en la ejecución final y Chrome durante el desarrollo.','Alcance limitado; Firefox y Safari pendientes.']],[1.25,3.3,2.2],font=9)
h('Métricas y alcance de los resultados')
p('La tasa de aprobación de los casos ejecutables fue 26/26, equivalente a 100 %. Este porcentaje no representa cobertura de código ni cobertura total de requisitos. La prueba de concurrencia obtuvo una sola reserva entre 20 solicitudes para el mismo intervalo. La revisión de interfaz comprobó ausencia de desbordamiento horizontal en los tamaños inspeccionados.')
h('Límites y siguientes validaciones')
p('Antes de usarlo con un negocio real se deben validar políticas de anticipación, cancelación y retrasos; medir con más datos y usuarios; realizar pruebas con clientes; revisar accesibilidad y probar navegadores adicionales. Para un despliegue público también hacen falta HTTPS, credenciales de base de datos adecuadas, respaldo y recuperación. Estas actividades amplían la evaluación local y no cambian la primera entrega.')
new('Conclusiones y referencias')
h('Conclusiones')
p('El proyecto convirtió el análisis de la primera entrega en un prototipo que permite reservar y administrar citas. La principal mejora es que la disponibilidad no depende únicamente de la pantalla: se calcula con jornadas, duración, descansos y bloqueos, y se vuelve a validar al guardar. La prueba de 20 solicitudes simultáneas respalda la consistencia de ese recorrido en el entorno evaluado.')
p('La cancelación y la reprogramación conservan los registros y sus movimientos. Esto permite explicar qué ocurrió con una cita y evita perder información cuando un cambio es rechazado. La configuración por trabajador y la asignación de servicios mantienen la relación entre lo ofrecido al cliente y la capacidad del negocio.')
p('Los resultados son favorables para la demostración académica. Todavía no permiten afirmar una reducción real de ausencias, satisfacción de clientes o disponibilidad continua. La siguiente validación debe realizarse con reglas acordadas con un encargado, usuarios representativos y un entorno de despliegue definido.')
h('Referencias')
refs=[
'[1] INGE00031. (s. f.). Producto Transversal. Ingeniería de Software. Documento de la actividad, páginas 4 a 6.',
'[2] Oracle. (s. f.). MySQL 8.0 Reference Manual. Locking Reads. https://dev.mysql.com/doc/refman/8.0/en/innodb-locking-reads.html',
'[3] Node.js. (s. f.). Crypto. crypto.scrypt. https://nodejs.org/api/crypto.html#cryptoscryptpassword-salt-keylen-options-callback',
'[4] ISO e IEC. (2023). ISO/IEC 25010:2023. Systems and software engineering. Systems and software Quality Requirements and Evaluation. Product quality model. Ficha pública del estándar. https://committee.iso.org/standard/78176.html',
'[5] Perea, E. (2026). Sistema web de gestión y control de citas para una barbería. Primera entrega. Documento local Sistema Web Barberia.docx.',
'[6] Proyecto Mala Vida Barbers. (2026). Código fuente, esquema y pruebas automatizadas. Archivos backend/src, database/schema.sql, frontend/src y backend/test. Ejecución de respaldo del 23 de septiembre de 2026.'
]
for r in refs:
 x=p(r);x.paragraph_format.space_after=Pt(8)
p('Fuentes web consultadas el 23 de septiembre de 2026. Los resultados numéricos proceden de la ejecución local del proyecto.')
h('Uso de herramientas de apoyo')
p('Se utilizó asistencia de inteligencia artificial para apoyar la implementación, la revisión del código y la preparación de esta documentación. Las afirmaciones sobre funcionamiento se vinculan con código y pruebas ejecutadas; no se presentan entrevistas ni resultados de usuarios que no se hayan realizado.')
new('Anexo A Trazabilidad funcional')
p('Esta matriz vincula los requisitos originales con la implementación y las pruebas del anexo B. La numeración RF se conserva. Las verificaciones corresponden al prototipo local.')
rf=[
('RF 01','Servicios','Alta, edición, precio, duración, descripción y estado activo.','T16 y T25'),
('RF 02','Barberos','Alta, edición, estado activo y servicios asignados.','T17 y T25'),
('RF 03','Jornadas y descansos','Periodos por trabajador y día; descanso entre periodos.','T06, T18 y T20'),
('RF 04','Disponibilidad','Intervalos para el servicio completo; excluye pasado, cruces y bloqueos.','T06, T10, T12 y T20'),
('RF 05','Registro de citas','Nombre y teléfono obligatorios; revalidación y transacción al guardar.','T07, T08, T09 y T10'),
('RF 06','Cancelación y reprogramación','Cambio atómico; rechazo conserva la cita; cancelación mantiene historial.','T11, T13 y T25'),
('RF 07','Agenda administrativa','Datos de cita, estados con texto, filtros por fecha y barbero y búsqueda.','T19 y T25'),
('RF 08','Autenticación','Inicio de sesión y rutas autorizadas; cambio de contraseña y cierre.','T05 y T26'),
('RF 09','Clientes','Registro relacionado con citas; búsqueda sin teléfono único.','T15 y T25'),
('RF 10','Estados','Cinco estados y transiciones restringidas. Reserva pública directamente confirmada.','T13 y T14; ajuste de flujo'),
('RF 11','Bloqueos','Intervalos especiales; citas afectadas señaladas sin eliminar registros.','T12, T20, T21 y T25'),
('RF 12','Búsqueda e historial','Nombre, teléfono o fecha; historial y movimientos de citas.','T11, T15, T19 y T25')]
table(['Requisito','Función','Implementación','Verificación'],rf,[.7,1.3,3.3,1.45],font=8.7)
p('Ajuste de RF-10: el flujo habitual pendiente a confirmada descrito en la primera entrega se simplifica para la reserva pública. El estado pendiente está admitido por el modelo y la lógica administrativa; su creación no tiene una pantalla específica ni una prueba dedicada de transición pendiente a confirmada. Debe explicarse como alcance real, no como escenario ya comprobado.')
new('Anexo A Requisitos no funcionales')
rnf=[
('RNF 01','Usabilidad','Formularios etiquetados, avisos, confirmaciones y recorrido automatizado.','Pendiente prueba con personas y registro de ayudas requeridas.'),
('RNF 02','Adaptación','Reserva y agenda inspeccionadas en escritorio y a 390 por 844 píxeles.','Sin desbordamiento en las vistas comprobadas; otros tamaños por revisar.'),
('RNF 03','Seguridad','scrypt, sesiones, validación, consultas parametrizadas y agenda protegida.','Controles básicos verificados; no auditoría integral.'),
('RNF 04','Rendimiento','20 consultas: mediana 16.5 ms y máximo 25.9 ms.','Cumple el umbral local de 2 s; carga reducida.'),
('RNF 05','Integridad','Bloqueo transaccional, rechazo de cruces y conservación tras rollback.','T08, T11 y T21 respaldan escenarios concurrentes.'),
('RNF 06','Mantenibilidad','Módulos de dominio, autenticación, datos, agenda y rutas; componentes compartidos.','Las rutas aún concentran parte de la coordinación.'),
('RNF 07','Compatibilidad','Edge en la corrida final y Chrome en recorridos previos.','Firefox y Safari no evaluados.')]
table(['Requisito','Aspecto','Evidencia','Límite'],rnf,[.7,1.1,2.6,2.35],font=9)
h('Relación con los casos previstos')
table(['Caso de la primera entrega','Prueba documentada','Qué demuestra'],[['CP 05 Reserva válida','T07 y T25','Reserva guardada, confirmación y retirada del horario disponible.'],['CP 06 Conflicto y concurrencia','T07 y T08','Diferencia entre cruce con una cita existente e intento simultáneo.'],['CP 07 Cambio rechazado','T11','Cita original intacta cuando el cambio falla.'],['CP 08 Permisos y bloqueos','T05 y T12','La agenda exige sesión y el bloqueo impide nuevas reservas.']],[1.8,1.45,3.5])
p('Un requisito puede tener código implementado y todavía requerir validación adicional. Se separa la evidencia técnica de la evaluación con usuarios, de la carga representativa y de la compatibilidad no comprobada. La aprobación de la suite no demuestra cumplimiento absoluto de todos los requisitos.')
cases=[
('T01','Fechas y marcas de tiempo','Aceptar fecha válida y rechazar imposibles o zona ambigua.','Aceptado y rechazado según el caso.'),
('T02','Cruce de intervalos','Permitir límites contiguos y detectar cruces.','Límites permitidos; cruces detectados.'),
('T03','Teléfonos e identificadores','Normalizar número y rechazar datos inválidos.','Validaciones cumplidas.'),
('T04','Aritmética civil','Sumar minutos al cambiar de día o año.','Cambio de fecha correcto.'),
('T05','Permisos y sesiones','Rechazar agenda sin sesión y comprobar cookie.','HTTP 401; HttpOnly y SameSite Strict.'),
('T06','Disponibilidad pública','Excluir descansos y espacios demasiado cortos.','Intervalos completos; sin datos de clientes.'),
('T07','Reserva y cruce parcial','Crear cita; rechazar cruce; aceptar límite adyacente.','HTTP 201, 409 y 201 respectivamente.'),
('T08','Concurrencia','Una sola cita ante 20 intentos simultáneos.','1 creada, 19 rechazadas; 1 registro.'),
('T09','Barberos diferentes','Permitir misma hora con otro trabajador.','Reserva aceptada.'),
('T10','Datos y fechas inválidas','Rechazar teléfono faltante, pasado, fecha imposible y fuera de jornada.','HTTP 400 o 409 según la regla.'),
('T11','Reprogramación','Conservar original al fallar y actualizar alternativa válida.','Inicio conservado; cambio válido con el mismo ID.'),
('T12','Bloqueos','Impedir reservas y señalar citas afectadas.','Rechazo y cita afectada conservada.'),
('T13','Cancelación','Liberar horario sin borrar ni reactivar una cita cerrada.','Historial conservado; reactivación rechazada.'),
('T14','Estados temporales','No completar ni marcar ausencia futura; completar pasada.','Futuras rechazadas; pasada completada.'),
('T15','Clientes e historial','Distinguir nombres con el mismo número.','Dos clientes e historial conservado.'),
('T16','Duración e inactividad','Conservar intervalos; ocultar y rechazar servicio inactivo.','Histórico sin cambios; reserva rechazada.'),
('T17','Barberos y servicios','Rechazar inactivo o servicio no asignado.','Rechazos y citas existentes conservadas.'),
('T18','Edición de jornadas','Rechazar cruces y señalar citas fuera de jornada.','Cruces rechazados; citas afectadas visibles.'),
('T19','Filtros y paginación','Combinar fecha, barbero, estado y búsqueda.','Registro esperado; página inválida rechazada.'),
('T20','Continuidad y varios días','Unir periodos contiguos y respetar bloqueo de varios días.','Intervalo continuo disponible; bloqueo sin horarios.'),
('T21','Reserva frente a bloqueo','Serializar operaciones concurrentes.','Cita rechazada o creada y señalada como afectada.'),
('T22','Tiempo de respuesta','Consultas locales en menos de 2 s.','Mediana 16.5 ms; máximo 25.9 ms.'),
('T23','Búsqueda parametrizada','No interpretar búsqueda como SQL.','Sin resultados ajenos; identificador inválido rechazado.'),
('T24','Inicialización repetida','No duplicar jornadas ni servicios.','Conteos conservados.'),
('T25','Recorrido de navegador','Reservar, administrar y revisar formularios y móvil.','Recorrido completado sin errores de página.'),
('T26','Contraseña y cierre','Invalidar sesiones y permitir nueva contraseña.','Sesión previa rechazada; nueva válida; cierre correcto.')]
for part in range(2):
 new('Anexo B Casos de prueba '+('de T01 a T13' if part==0 else 'de T14 a T26'))
 p('Resultado de todos los casos de esta tabla: aprobado en la ejecución registrada. Los identificadores T se asignan para esta documentación y se vinculan con las pruebas automatizadas del proyecto.')
 table(['ID','Escenario','Resultado esperado','Resultado observado'],cases[part*13:part*13+13],[.48,1.45,2.42,2.4],font=8.8)
 p('Fuente: backend/test/domain.test.js y backend/test/integration.test.js. Los primeros cuatro casos evalúan reglas aisladas; los demás utilizan integración y, en T25, navegador real.')
new('Anexo C Evidencia de la interfaz pública')
p('La captura procede del recorrido automatizado de reserva. Servicio, barbero y horario se seleccionaron en la interfaz; nombre y teléfono pertenecen al conjunto ficticio de pruebas.')
pic(ROOT/'backend/test-results/reserva-escritorio.png',6.5);cap('Figura C1. Reserva en escritorio durante T25.')
h('Qué se puede explicar con esta imagen')
bullet('El cliente conoce duración y precio antes de confirmar.')
bullet('La selección de servicio determina los barberos que pueden realizarlo.')
bullet('Los horarios se calculan desde la configuración y las citas, no son un campo libre.')
bullet('La confirmación definitiva ocurre después de validar y guardar en el servidor.')
p('Para las diapositivas puede reutilizarse la imagen completa o una captura del formulario. La demostración debe mostrar además la respuesta de confirmación; una imagen por sí sola no demuestra la persistencia de la cita.')
new('Anexo C Evidencia de la administración')
pic(ROOT/'backend/test-results/agenda-escritorio.png',6.5);cap('Figura C2. Agenda después de reprogramar y cancelar una cita durante T25. Datos ficticios.')
h('Qué demuestra el recorrido')
p('El administrador inició sesión, encontró una cita mediante búsqueda, modificó su horario y confirmó la cancelación. La tarjeta permanece visible con el estado Cancelada. Desde Clientes se consultó el historial y se verificó ese mismo resultado.')
h('Comprobación móvil')
p('La prueba cambió el tamaño de pantalla a 390 por 844 píxeles y comprobó que el ancho del documento no superara el ancho de la ventana. Se revisaron reserva y agenda. Esto respalda la adaptación en ese tamaño específico, sin representar una prueba en todos los teléfonos.')
h('Archivos de evidencia disponibles')
p('Las capturas originales están en backend/test-results: reserva-escritorio.png, agenda-escritorio.png, reserva-movil.png y agenda-movil.png. Contienen datos de prueba y pueden emplearse como apoyo para la exposición.')
new('Anexo D Instalación y operación')
h('Preparación del entorno')
p('Se requiere Node.js 22.12 o posterior, Docker Desktop en ejecución y los puertos locales 3001 y 3306 disponibles. Desde la carpeta principal, ejecutar estos pasos:')
for cmd,desc in [('npm run setup','Instala las dependencias del servidor y la interfaz.'),('npm run db','Inicia MySQL y espera a que esté disponible.'),('npm run build','Genera la interfaz que entrega el servidor.'),('npm start','Inicia la aplicación en http://127.0.0.1:3001.')]:p(cmd+'  -  '+desc)
p('La primera ejecución crea un administrador y registra sus credenciales en backend/.admin-access.txt. Ese archivo es privado y está excluido de Git. No debe incluirse en el informe, las diapositivas ni un repositorio compartido. La contraseña se cambia desde Administración y Mi cuenta.')
h('Operación cotidiana')
p('El administrador mantiene los servicios y sus duraciones, asigna servicios a cada barbero y configura la jornada. Para una ausencia excepcional crea un bloqueo. Si el sistema señala citas afectadas, debe resolverlas mediante reprogramación o cancelación; crear un bloqueo no borra las citas.')
p('El cliente elige servicio, barbero, fecha y horario disponible; después registra nombre y teléfono y confirma. Para solicitar cambios se comunica con el administrador. El número de reserva permite localizar la cita, pero no funciona como credencial pública para acceder a datos de clientes.')
h('Repetir las pruebas')
p('npm test ejecuta los casos del servidor. Para el recorrido visual, compilar primero la interfaz, definir BROWSER_CHANNEL=msedge o chrome y ejecutar npm run test:browser. En PowerShell la variable se define con $env:BROWSER_CHANNEL = "msedge". Las pruebas crean y eliminan una base temporal propia.')
h('Detención y recuperación')
p('Para detener MySQL sin borrar datos se utiliza docker compose stop. El volumen db_data conserva los registros entre reinicios. Si no carga el catálogo, revisar que Docker y el servidor estén iniciados. Si el acceso falla, comprobar las credenciales vigentes y respetar mayúsculas; después de cambiar la contraseña, el archivo inicial deja de representar la contraseña actual.')
h('Desarrollo y configuración')
p('npm run dev inicia servidor e interfaz de desarrollo en el puerto 5173. No debe ejecutarse a la vez que npm start porque ambos intentan iniciar la API. backend/.env.example contiene las variables opcionales de conexión. No es necesario modificar la primera entrega para ejecutar el sistema.')
new('Anexo E Contenido para las diapositivas 1 a 4')
slides=[
('1 Presentación del proyecto','Sistema web de citas para Mala Vida Barbers. Segunda entrega de Ingeniería de Software. Diseño, implementación y evaluación.','El proyecto parte de una necesidad concreta: mantener una agenda compartida que no acepte reservas incompatibles. En esta fase se implementó el prototipo y se comprobó su comportamiento con pruebas.','Portada y nombre del sistema.'),
('2 Problema y objetivo','Información dispersa entre mensajes y anotaciones. Riesgo de cruces y pérdida de cambios. Objetivo: centralizar reservas e historial.','Una cita ocupa un intervalo completo. No basta con verificar su hora inicial; también hay que considerar el tiempo del servicio y la disponibilidad del trabajador.','Ejemplo de dos citas que se cruzan.'),
('3 Alcance y usuarios','Cliente sin cuenta. Administrador autenticado. Servicios, barberos, jornadas, bloqueos, citas e historial.','El cliente reserva y el administrador mantiene la operación. Se conserva el alcance de la primera entrega: pagos, inventario y recordatorios automáticos quedan fuera.','Funciones principales o pestañas del panel.'),
('4 Arquitectura','React y Tailwind en el navegador. API con Node.js y Express. MySQL para datos y relaciones.','El navegador solicita operaciones al servidor. Las validaciones importantes se ejecutan allí, antes de modificar la base; por eso no dependen de ocultar botones.','Figura 1 de este documento.'),
('5 Modelo de datos','Citas vinculadas con cliente, servicio y barbero. Inicio y fin guardados. Jornadas, bloqueos y movimientos separados.','La duración reservada se conserva aunque después cambie el catálogo. El historial permite consultar citas canceladas y los cambios realizados.','Figura 3 simplificada.'),
('6 Regla de disponibilidad','Duración completa dentro de la jornada. Sin descansos, bloqueos ni citas activas superpuestas. Validación repetida al guardar.','La disponibilidad puede cambiar mientras alguien llena el formulario. Por eso la confirmación vuelve a revisar el intervalo dentro de una transacción.','Figura 2 y ejemplo de intervalo.'),
('7 Demostración del prototipo','Reservar una cita. Localizarla en administración. Reprogramar y cancelar conservando el historial.','Voy a mostrar el recorrido del cliente y después el del administrador. La cancelación libera el horario, pero el registro continúa visible con su estado.','Demostración en vivo; capturas C1 y C2 como respaldo.'),
('8 Pruebas funcionales','26 casos ejecutables aprobados. El registro suma 27 al incluir el grupo de integración. Sin fallos en la corrida de respaldo.','Las pruebas usan solicitudes reales y una base independiente. Incluyen validaciones, permisos, cambios y navegador. El porcentaje aprobado no equivale a cobertura total de código.','Resumen de la página 8.'),
('9 Concurrencia y fiabilidad','20 intentos simultáneos. 1 reserva creada. 19 rechazos por conflicto. Reprogramación fallida conserva la cita.','Las solicitudes compiten por el mismo espacio y solo una lo obtiene. El control se aplica en MySQL y no solo en la interfaz.','Tres cifras y explicación del bloqueo.'),
('10 Evaluación de calidad','Disponibilidad local: mediana 16.5 ms y máximo 25.9 ms. Acceso administrativo protegido. Vista móvil comprobada.','Los tiempos corresponden a veinte consultas secuenciales locales con pocos datos. La evaluación sigue criterios medibles, pero no constituye una certificación.','Tabla breve de calidad.'),
('11 Límites y mejoras','Faltan pruebas con usuarios y mayor carga. Firefox y Safari pendientes. Validar políticas con el negocio.','El prototipo permite una demostración académica. Antes de un uso real se necesitan reglas acordadas, evaluación adicional y un entorno de despliegue seguro.','Tres pendientes verificables.'),
('12 Cierre','Agenda centralizada. Cambios con historial. Evidencia reproducible de pruebas.','El resultado principal es una reserva consistente y administrable. La ingeniería permite vincular lo solicitado, lo implementado y lo comprobado, dejando visibles sus límites.','Conclusión y preguntas.')]
for i,slide in enumerate(slides):
 if i in [4,8]:new('Anexo E Contenido para las diapositivas '+('5 a 8' if i==4 else '9 a 12'))
 title,screen,notes,visual=slide
 h(title);p('Texto para la diapositiva: '+screen);p('Guion oral: '+notes);x=p('Apoyo visual: '+visual);x.paragraph_format.space_after=Pt(14)
new('Anexo F Guion de demostración y preguntas')
p('Duración sugerida: cinco a siete minutos. Preparar antes un día futuro con jornada y un barbero activo que ofrezca el servicio. Usar nombres y teléfonos ficticios.')
table(['Paso','Acción','Resultado que se explica'],[['1','Consultar servicio, trabajador y fecha.','Solo intervalos donde cabe el servicio.'],['2','Reservar con datos de prueba.','Número de cita y confirmación.'],['3','Entrar como administrador y filtrar.','Los datos están en una vista protegida.'],['4','Reprogramar a otro horario libre.','Misma cita, nuevo intervalo y movimiento registrado.'],['5','Crear un bloqueo que afecte esa cita.','Se señala el conflicto sin borrar la reservación.'],['6','Resolver cancelando o reprogramando.','Historial conservado y disponibilidad actualizada.'],['7','Mostrar el resultado automatizado.','20 solicitudes: una aceptación y 19 rechazos.']],[.5,2.8,3.45])
h('Por qué no basta con validar desde React')
p('Dos clientes pueden consultar el mismo horario antes de que alguno guarde. La decisión final debe tomarse en el servidor dentro de una operación protegida.')
h('Por qué conservar una cita cancelada')
p('Su estado forma parte del historial. Cancelar libera capacidad; borrar perdería la información sobre lo que ocurrió.')
h('Qué demuestra el resultado de las pruebas')
p('Los casos ejecutados cumplieron las expectativas en ese entorno. No demuestra ausencia de todos los defectos, cobertura completa, satisfacción de clientes ni desempeño con otra carga.')
h('Qué cambia frente a la primera entrega')
p('La propuesta se convirtió en código y se fijaron supuestos operativos. La reserva pública se confirma directamente; pendiente permanece en el modelo. Las políticas siguen necesitando validación con el negocio y las herramientas planificadas no se presentan como utilizadas si no existe evidencia.')
new('Anexo G Reparto del material y cierre de entrega')
table(['Destino','Contenido que puede tomarse de este documento'],[['Informe final','Páginas 1 a 10. Completar portada y revisar las indicaciones del docente.'],['Anexos técnicos','Páginas 11 a 16. Matrices, casos y capturas de respaldo.'],['Manual e instalación','Página 17 y README del proyecto.'],['Presentación','Páginas 18 a 20. Doce diapositivas con texto, guion y apoyo visual.'],['Demostración','Página 21. Recorrido y respuestas para la exposición.']],[1.65,5.1])
h('Reparto sugerido si participan varias personas')
p('Una persona puede explicar problema, alcance y arquitectura; otra presentar datos y reglas; otra realizar la demostración; y otra exponer pruebas, calidad y límites. Si la exposición es individual, conservar ese orden. Los nombres y tiempos se asignan según los participantes reales.')
h('Lista de cierre')
for item in ['Completar docente, institución y grupo; confirmar el nombre del alumno.', 'Revisar si los anexos se entregan juntos o separados del informe de 10 páginas.', 'Copiar los textos de las diapositivas y seleccionar capturas legibles al proyectar.', 'Habilitar un repositorio accesible y añadir su enlace. No hay un remoto verificado en la carpeta revisada.', 'Excluir contraseñas, archivos .env y .admin-access.txt del repositorio y del material compartido.', 'Probar la instalación desde el equipo de exposición y preparar capturas de respaldo.', 'Validar políticas con un encargado si la actividad exige datos reales; conservar evidencia antes de afirmar que se realizó.', 'Entregar en la plataforma y fecha indicadas por el docente.']:bullet(item)
h('Afirmaciones que requieren evidencia adicional')
p('No atribuir al sistema una certificación, pruebas con clientes, reducción medida de ausencias, cobertura de código del 100 %, funcionamiento comprobado en Safari o Firefox, ni publicación en un repositorio que todavía no sea accesible. La precisión de estas afirmaciones forma parte de la calidad de la entrega.')
for para in D.paragraphs:para.paragraph_format.widow_control=True
for element in [D.element, D.styles.element]:
 for border in element.xpath('.//w:pBdr'):border.getparent().remove(border)
D.save(OUT/'Documentacion_segunda_entrega_barberia.docx')
print('Paginas planificadas',page)
print(OUT/'Documentacion_segunda_entrega_barberia.docx')


