from docx import Document
from pathlib import Path
from docx.oxml.ns import qn
D=Document('Documentacion_segunda_entrega_barberia.docx')
changes={
8:'Este informe explica cómo funciona el sistema de citas de la barbería y qué resultados obtuvo en las pruebas. Retoma los requisitos de la primera entrega para describir el diseño y la implementación de esta segunda etapa. También sirve como base para presentar y demostrar el proyecto.',
10:'El sistema permite reservar citas, administrar horarios y consultar el historial de cambios. En la prueba local aprobaron los 26 casos ejecutados. El registro muestra 27 resultados porque también cuenta la aprobación del grupo de integración. No hubo fallos.',
14:'Para dar una cita, la barbería necesita saber quién puede atender y cuánto tiempo requiere el servicio. Si las reservas quedan repartidas entre mensajes y notas, un mismo horario puede asignarse dos veces. También es fácil perder de vista un cambio. La primera entrega describió este problema, las personas que usarían el sistema y sus requisitos.',
15:'En esta etapa, la propuesta tomó forma en un sitio web. El cliente puede consultar los servicios, elegir un horario y reservar sin abrir una cuenta. El administrador sí debe iniciar sesión. Desde su panel revisa la agenda, modifica los catálogos y atiende cancelaciones, cambios de horario o ausencias.',
16:'La prioridad fue comprobar que las reglas funcionaran al guardar una cita. Por ejemplo, si dos personas eligen el mismo horario, el sistema debe aceptar solo una reserva. Las pruebas de esta entrega revisan ese caso y otras operaciones de la agenda.',
18:'El sistema incluye servicios, barberos, clientes y citas, además de jornadas, descansos, bloqueos y acceso administrativo. Quedan fuera los pagos en línea, el inventario, la nómina, los recordatorios automáticos y las cuentas individuales de los barberos. Los horarios y las políticas que usa la demostración son supuestos académicos. Todavía falta acordarlos con un negocio real.',
22:'El sistema sigue una arquitectura cliente-servidor. En el navegador, React muestra los formularios y la agenda. Cuando una persona reserva o modifica una cita, la solicitud llega a la API, que es la vía de comunicación con el servidor. Allí, Node.js y Express revisan las reglas y consultan MySQL. El navegador nunca accede por su cuenta a la base de datos.',
24:'Figura 1. Comunicación entre la pantalla, el servidor y la base de datos. La interfaz intercambia datos por HTTP en formato JSON. El servidor usa consultas SQL con parámetros.',
27:'El código del servidor está dividido según su tarea: revisar datos, gestionar el acceso, consultar la base o aplicar las reglas de la agenda. Cada pantalla administrativa también tiene su propio componente. Las pantallas comparten funciones para enviar solicitudes y mostrar mensajes o diálogos. Así, una corrección en esas funciones puede aprovecharse en varias pantallas.',
29:'Cualquier visitante puede consultar servicios, barberos y horarios libres, pero no los nombres ni los teléfonos de otros clientes. Para ver la agenda o hacer cambios se necesita una sesión vigente. Las contraseñas se procesan con scrypt y una sal aleatoria, en lugar de guardarse como texto. La cookie de sesión usa HttpOnly y SameSite Strict. La referencia [3] explica la función scrypt de Node.js.',
30:'Al ejecutar la versión local, la interfaz y la API usan la misma dirección. Durante el desarrollo, Vite envía las solicitudes de la API al servidor. Esto permite editar las pantallas mientras las reglas de las citas siguen a cargo del servidor.',
33:'Un horario puede estar libre al consultarlo y ocuparse antes de confirmar. Por eso, el servidor lo revisa otra vez al guardar. Esa revisión y el registro de la cita forman parte de una transacción: si algo falla, la operación se revierte.',
35:'Figura 2. Pasos para reservar una cita. Si la validación falla, el proceso sigue la ruta inferior izquierda y devuelve un aviso al cliente.',
37:'Dos citas se cruzan cuando inicio_nueva < fin_existente y fin_nueva > inicio_existente. La regla se aplica al mismo barbero y a citas pendientes o confirmadas. En cambio, una cita que termina a las 10:40 permite que otra empiece a esa hora. Además, todo el servicio debe caber en la jornada, sin ocupar descansos ni periodos bloqueados.',
39:'Para evitar que dos solicitudes guarden el mismo horario, el servidor usa SELECT FOR UPDATE sobre una fila de control [2]. Mientras una operación revisa y guarda los datos, las demás esperan. Al continuar, vuelven a comprobar si el horario sigue libre. El bloqueo termina al confirmar o revertir la transacción. Este mecanismo ordena las escrituras, aunque también puede limitar la capacidad del sistema cuando aumente la demanda.',
41:'Al reprogramar, primero se revisa el nuevo horario. Si no es válido, la cita conserva sus datos. La cancelación libera el espacio y mantiene el historial. Una cita confirmada puede marcarse como completada después de terminar, o como ausencia después de su inicio. En la reserva pública, las citas quedan confirmadas desde el principio. El modelo admite el estado pendiente y sus cambios, pero el formulario público no permite crearlo. Este comportamiento ajusta el recorrido previsto en la primera entrega.',
44:'Las tablas separan los catálogos de las citas y de los datos de acceso. Cada cita se vincula con un servicio y un barbero. Los registros nuevos también se vinculan con un cliente. Las claves foráneas mantienen estas relaciones. Cuando un servicio o barbero deja de estar disponible, se marca como inactivo para conservar los datos anteriores.',
46:'Figura 3. Relaciones de la base de datos. Las flechas indican qué registros dependen de otros. Por ejemplo, un cliente puede tener varias citas, y cada cita puede tener varios movimientos.',
48:'Al reservar, el sistema guarda la hora de inicio y calcula la de fin. Si después cambia la duración del servicio, las citas anteriores conservan su horario. También guardan el nombre y el teléfono recibidos en la reserva. Un teléfono puede pertenecer a más de una persona. Si nombre y teléfono coinciden exactamente, se reutiliza el cliente registrado; cuando dos personas comparten ambos datos, el administrador debe revisar el caso.',
49:'Un barbero puede tener varios periodos de trabajo en el mismo día. El espacio entre ellos cuenta como descanso. Si un periodo termina justo cuando empieza otro, el sistema los trata como una jornada continua. Los bloqueos tienen inicio y fin, y pueden durar varios días. Cuando afectan una cita ya guardada, aparece un aviso para que el administrador la resuelva. La cita no se borra.',
50:'Otras tablas guardan administradores y sesiones (admins y admin_sessions), cambios de las citas (appointment_events), el bloqueo de agenda (schedule_mutex) y las actualizaciones de estructura (migrations). Los índices ayudan a consultar la agenda y buscar registros. Al iniciar, el sistema adapta la base anterior y vincula sus citas con clientes sin borrar la información existente.',
53:'La pantalla conserva el fondo oscuro y los detalles dorados de Mala Vida Barbers. El cliente encuentra allí los servicios y elige barbero, fecha y hora. Luego escribe sus datos de contacto y confirma. Al terminar, recibe el número de reserva junto con el horario asignado.',
55:'Figura 4. Pantalla de reserva durante una prueba automática. El nombre y el teléfono que aparecen son datos de prueba.',
57:'El panel tiene apartados para agenda, clientes, servicios, barberos, jornadas, bloqueos y cuenta. Para encontrar una cita se puede filtrar por fecha, barbero, estado, nombre o teléfono. Antes de cancelar o cambiar el estado, el sistema pide confirmar la acción. Después informa si pudo guardarla o si hubo un conflicto. También avisa cuando un bloqueo o un cambio de jornada afecta citas existentes.',
58:'En el teléfono, los formularios y las tarjetas aparecen en una sola columna. La revisión a 390 por 844 píxeles no mostró desplazamiento horizontal. Los campos tienen etiquetas, el foco del teclado es visible y los estados se leen en texto, además de distinguirse por color. Aún falta evaluar la accesibilidad con más detalle y observar cómo usan el sistema personas reales.',
62:'La interfaz usa React, Vite y Tailwind CSS. El servidor está escrito en JavaScript con Node.js y Express, y se conecta a MySQL mediante mysql2. Docker guarda la base en un volumen para conservarla entre reinicios. Las funciones del proyecto se agrupan en los apartados que muestra la siguiente tabla.',
64:'El plan de la primera entrega incluía cuatro iteraciones de dos semanas y proponía usar Trello, Figma y Visual Paradigm. Para esta entrega se cuenta con el código, los diagramas del informe y los registros de pruebas. Esos materiales permiten revisar el trabajo realizado, aunque no acreditan fechas de cada iteración, entrevistas ni el uso de todas las herramientas propuestas.',
66:'Al iniciar, el sistema crea las tablas que faltan y actualiza la estructura anterior. Deja un registro de esa actualización para no repetir la carga. Antes de guardar, revisa identificadores, fechas, teléfono y duración, y usa parámetros en las consultas. Los cambios de agenda comparten el mismo control de transacciones. Una prueba comprobó que iniciar de nuevo no duplicara servicios ni jornadas.',
68:'La demostración usa la hora de Ciudad de México. El horario inicial va de lunes a sábado, de 10:00 a 14:00 y de 15:00 a 20:00, aunque puede cambiarse para cada barbero. El cliente puede reservar hasta con 90 días de anticipación y ve opciones cada 10 minutos. Para dar de alta un barbero hay que asignarle una jornada. También se debe indicar qué barberos realizan cada servicio nuevo.',
69:'El sistema se abre localmente en http://127.0.0.1:3001. El archivo README del proyecto explica cómo instalarlo y ponerlo en marcha. Falta publicar el repositorio remoto y añadir su enlace a la entrega. La dirección local solo sirve en el equipo donde se ejecuta la aplicación.',
73:'Las pruebas que respaldan este informe se ejecutaron el 23 de septiembre de 2026 en Windows, con Node.js 24.19.0, MySQL 8.0.45 y Microsoft Edge 153.0.4234.48. Usaron datos ficticios en una base temporal y enviaron solicitudes HTTP al servidor. Playwright 1.58.2 automatizó las acciones en el navegador. Al terminar, se borró únicamente la base creada para esas pruebas.',
74:'Aprobaron cuatro casos de reglas básicas y 22 de integración: 26 casos en total. El registro cuenta además la aprobación del grupo de integración, por lo que muestra 27 resultados. La tabla resume las comprobaciones principales. Los casos y sus resultados pueden consultarse en los archivos de pruebas del proyecto.',
77:'Se midieron 20 consultas de disponibilidad, una después de otra, con conexión local y una base pequeña. La mediana fue de 16.5 ms y el valor más alto, de 25.9 ms. Ambos quedaron por debajo de los 2 segundos establecidos en RNF-04. Estos tiempos corresponden a la respuesta de la API. Falta medir la carga completa de la página y el comportamiento con más usuarios.',
79:'Durante las pruebas se corrigieron las etiquetas que permiten identificar los selectores y el manejo de periodos de trabajo consecutivos. También terminaron sin errores ni advertencias la compilación y el análisis estático de la interfaz. Los casos revisados funcionaron como se esperaba, aunque pueden existir fallos en situaciones que todavía no se han probado.',
82:'La revisión de calidad usa ISO/IEC 25010:2023 como referencia general [4]. La tabla reúne los aspectos que se pudieron revisar en este prototipo y señala lo que falta comprobar. Esta evaluación local no equivale a una certificación ni a una revisión completa de la norma.',
85:'Los 26 casos ejecutados aprobaron, así que la tasa de aprobación fue del 100 %. Ese dato describe esta ejecución y no indica qué porcentaje del código o de los requisitos quedó cubierto. En el intento simultáneo de 20 reservas para el mismo horario, solo una se guardó. Las pantallas revisadas también se mantuvieron dentro del ancho disponible.',
87:'Antes de usar el sistema en una barbería, hace falta acordar con el encargado las reglas de anticipación, cancelación y retrasos. Después conviene probarlo con clientes, más registros y otros navegadores, además de revisar su accesibilidad. Publicarlo requiere HTTPS, credenciales de base de datos adecuadas y un procedimiento de respaldo y recuperación. Estos pasos corresponden a la continuación del proyecto y no modifican la primera entrega.',
91:'La segunda etapa deja un prototipo con el que se pueden reservar y administrar citas. La disponibilidad considera la duración del servicio y los horarios del barbero, incluidos sus descansos y bloqueos. Al guardar, el servidor vuelve a revisar esas condiciones para evitar aceptar un espacio que ya se ocupó.',
92:'Los cambios de horario y las cancelaciones quedan en el historial. Esto permite consultar qué pasó con una cita, incluso si un intento de cambio fue rechazado. Además, la configuración de cada barbero determina qué servicios puede atender y en qué horarios aparecen disponibles.',
93:'Las pruebas realizadas respaldan el uso del prototipo en una demostración académica. Para conocer su utilidad en la operación diaria, todavía hay que probarlo con un encargado y con clientes. Por ahora no se ha medido una reducción de ausencias, la satisfacción de quienes lo usan ni su funcionamiento continuo en un entorno público.',
101:'Se usó inteligencia artificial como apoyo para implementar y revisar el código, y para preparar y revisar la redacción del informe.',
102:'Los resultados descritos se basan en el código y en las pruebas del proyecto. No se realizaron entrevistas ni pruebas con clientes reales.'
}
def replace(p,t):
 runs=p.runs
 if runs:
  runs[0].text=t
  for r in runs[1:]:r.text=''
 else:p.add_run(t)
for i,t in changes.items():replace(D.paragraphs[i],t)
# Desarrollar etiquetas comprimidas sin cambiar los resultados.
table_changes={
(2,1,1):'Calcula horarios válidos y rechaza cruces, incluso ante solicitudes simultáneas.',
(2,2,1):'Permite iniciar sesión, filtrar citas, cambiar su horario, cancelarlas y actualizar su estado.',
(2,3,1):'Permite asignar servicios a los barberos y definir sus jornadas y bloqueos.',
(2,4,1):'Conserva clientes e historial e incluye pruebas de la API y del navegador.',
(3,2,1):'El cambio se rechazó y la cita conservó su identificador y su hora de inicio.',
(3,3,1):'La cita siguió en el historial y el horario volvió a estar libre.',
(3,4,1):'El servidor rechazó las consultas sin sesión con la respuesta HTTP 401.',
(3,5,1):'El sistema excluyó los periodos ocupados y avisó de las citas afectadas.',
(3,7,1):'27 resultados aprobados, ninguno fallido ni omitido. Duración aproximada: 14.32 segundos.',
(4,1,1):'Las pruebas revisan las reservas y los cambios de las citas.',
(4,1,2):'Las funciones principales están listas. La reserva pública confirma directamente.',
(4,2,1):'Se probaron cruces, cambios rechazados, solicitudes simultáneas e historial.',
(4,2,2):'Los casos aprobaron. Falta probar fallos prolongados.',
(4,3,1):'En 20 consultas locales: mediana de 16.5 ms y máximo de 25.9 ms.',
(4,3,2):'Se cumplió la meta local de 2 s. Falta probar con mayor carga.',
(4,4,1):'Exige sesión y usa scrypt, opciones de protección de cookies y consultas con parámetros.',
(4,4,2):'Se revisaron controles básicos. Falta una prueba de penetración.',
(4,5,1):'Incluye etiquetas, avisos y confirmaciones. La vista móvil cabe en la pantalla.',
(4,5,2):'La revisión técnica fue favorable. Falta probar con usuarios.',
(4,6,1):'El código separa módulos, comparte componentes y permite repetir pruebas.',
(4,6,2):'La estructura facilita revisar el código. No se midió su cobertura.',
(4,7,1):'Se usó Edge en la prueba final y Chrome durante el desarrollo.',
(4,7,2):'Falta comprobar Firefox y Safari.'
}
for (ti,ri,ci),t in table_changes.items():
 cell=D.tables[ti].cell(ri,ci);replace(cell.paragraphs[0],t)
 for p in cell.paragraphs[1:]:replace(p,'')
# Retirar únicamente los párrafos vacíos finales que no contienen ilustraciones.
while D.paragraphs and not D.paragraphs[-1].text.strip() and not D.paragraphs[-1]._p.xpath('.//w:drawing'):
 el=D.paragraphs[-1]._p;el.getparent().remove(el)
for el in [D.element,D.styles.element]:
 for border in el.xpath('.//w:pBdr'):border.getparent().remove(border)
out=Path('output/documentos/Documentacion_segunda_entrega_barberia_revisada.docx')
out.parent.mkdir(parents=True,exist_ok=True);D.save(out)
print(f'{len(changes)} párrafos y {len(table_changes)} celdas revisados. {len(D.inline_shapes)} imágenes conservadas.')
